"""기관-학생 매칭 결과 xlsx → 직무별 배정 수 CSV (job.final_assigned 입력).

    python matching_counts.py "<매칭 결과.xlsx>" --out out/final_assigned.csv

- 읽는 칸은 머리글이 '지원 기관명'·'지원 직무'인 두 열뿐이다(2026-2 파일은 6·7열). 이름·학과·연락처 열은 읽지도 출력하지도 않는다.
- 결과는 (기관, 직무, 배정) 집계만. 학과별로 나누지 않는다(ADR-0004 — 칸당 1~2명이라 식별 위험).
- 시트 '공지용', 머리글 8행(2026-2 파일 기준). 다른 학기 파일이면 --sheet·--header-row를 맞춘다.
"""
import argparse, collections, csv, pathlib, re

import openpyxl

HEAD_INSTITUTION, HEAD_JOB = "지원 기관명", "지원 직무"


def clean(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def read_counts(path, sheet="공지용", header_row=8):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    # 머리글 행에서 두 열의 위치만 찾는다(머리글 글자만 본다)
    heads = {clean(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1)}
    if HEAD_INSTITUTION not in heads or HEAD_JOB not in heads:
        raise SystemExit(f"{header_row}행에 '{HEAD_INSTITUTION}'·'{HEAD_JOB}' 머리글이 없습니다 — --sheet·--header-row를 확인하세요")
    col_inst, col_job = heads[HEAD_INSTITUTION], heads[HEAD_JOB]
    grid = {}
    for col in (col_inst, col_job):
        for r in range(header_row + 1, ws.max_row + 1):
            grid[(r, col)] = ws.cell(r, col).value
    # 기관명 칸이 세로로 병합된 경우 위 값으로 채운다. 병합의 왼쪽 위 칸이 두 열 중 하나일 때만 —
    # 다른 열에서 시작한 병합이면 그 값은 개인 정보 열의 것일 수 있어 쓰지 않는다
    for rng in ws.merged_cells.ranges:
        if rng.max_row <= header_row or rng.min_col not in (col_inst, col_job) or rng.max_col > max(col_inst, col_job):
            continue
        top = ws.cell(rng.min_row, rng.min_col).value
        for col in range(rng.min_col, rng.max_col + 1):
            for r in range(max(rng.min_row, header_row + 1), rng.max_row + 1):
                grid[(r, col)] = top
    counts = collections.Counter()
    for r in range(header_row + 1, ws.max_row + 1):
        inst, job = clean(grid[(r, col_inst)]), clean(grid[(r, col_job)])
        if inst or job:
            counts[(inst, job)] += 1
    return counts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="공지용")
    ap.add_argument("--header-row", type=int, default=8)
    ap.add_argument("--out", default=str(pathlib.Path(__file__).parent / "out" / "final_assigned.csv"))
    a = ap.parse_args()
    counts = read_counts(a.xlsx, a.sheet, a.header_row)
    out = pathlib.Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["기관", "직무", "배정"])
        for (inst, job), n in sorted(counts.items()):
            w.writerow([inst, job, n])
    print(f"직무 {len(counts)}개, 배정 합 {sum(counts.values())}명 → {out}")


if __name__ == "__main__":
    main()
