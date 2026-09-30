"""참여기관 리스트 xlsx 파서 (E5·E6 공용).

'(전체)2026-2학기 참여기관' 시트: 3행 헤더, 4~43행 직무 40개, 병합 셀 65개(기관 단위 칸이 세로로 병합).
병합 셀은 왼쪽 위 값으로 채워서 직무 행마다 기관 정보가 들어가게 만든다.
"""
import re
import openpyxl

SHEET = "(전체)2026-2학기 참여기관"
HEADER_ROW = 3
COLS = ["순번", "현장실습참여여부", "취업연계참여여부", "기관명", "기간요일시간", "부서직무",
        "요구역량", "모집인원", "근로지주소", "선호전공학년", "실습지원비", "비고"]


def clean(v):
    if v is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def read_jobs(path, sheet=SHEET):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    grid = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, max_row=ws.max_row, max_col=len(COLS)):
        for c in row:
            grid[(c.row, c.column)] = c.value
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if r > HEADER_ROW:
                    grid[(r, col)] = top
    jobs = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        row = {name: clean(grid.get((r, i + 1))) for i, name in enumerate(COLS)}
        if not row["부서직무"] and not row["기관명"]:
            continue
        row["기관"] = row["기관명"].split("\n")[0].strip()
        row["행"] = r
        jobs.append(row)
    return jobs


# 수기·리스트에서 같은 기관을 가리키는 표준 키
CANON = [("소서", ["소서", "원오세븐"]), ("오뷔엘알", ["오뷔엘알", "올리비아로렌"]), ("델타텍코리아", ["델타텍"]),
         ("위아프렌즈", ["위아프렌즈"]), ("챔프스터디", ["챔프스터디", "해커스"]), ("더에스엠씨", ["더에스엠씨"]),
         ("빌딩닥터엔지니어링", ["빌딩닥터"]), ("스미스", ["스미스"])]


def canon(name, department=""):
    s = re.sub(r"\(주\)|㈜|주식회사|\s", "", name or "")
    if "세정" in s and "OL" in (department or "").upper():
        return "오뷔엘알"          # 수기의 '세정 OL디자인팀' = 현재 오뷔엘알
    for key, aliases in CANON:
        if any(a in s for a in aliases):
            return key
    return s
